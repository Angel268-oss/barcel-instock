"""
convertir.py — Lee el Excel de Barcel y genera datos.json para la página web.
Uso: python convertir.py "Seguimiento indicadores Barcel.xlsm"
"""

import sys
import json
import re
from datetime import datetime, date
from openpyxl import load_workbook

REGION = "10ZAI"

REINC_ORDER = {"REC_CRITICA": 0, "REC_ALTA": 1, "REINCIDENTE": 2, "SIN_REINC": 3}


def plaza_corta(plaza_str):
    """'10ZAI Monclova' → 'Monclova'"""
    return plaza_str.replace(f"{REGION} ", "").strip()


def get_fecha(wb):
    """Intenta leer la fecha de la hoja ONEPAGE o la de hoy."""
    try:
        ws = wb["ONEPAGE"]
        for row in ws.iter_rows(max_row=10, values_only=True):
            for cell in row:
                if isinstance(cell, (datetime, date)):
                    if isinstance(cell, datetime):
                        return cell.date()
                    return cell
    except Exception:
        pass
    return date.today()


def leer_instock(wb):
    """Lee hoja INSTOCK DETALLE → dict CR_TIENDA → {instock, clase, rango}"""
    ws = wb["INSTOCK DETALLE"]
    data = {}
    for row in ws.iter_rows(values_only=True):
        plaza = row[3]
        if not plaza or REGION not in str(plaza):
            continue
        cr = str(row[5]).strip()
        try:
            instock = float(row[6]) * 100
        except (TypeError, ValueError):
            instock = None
        clase = str(row[7]).strip() if row[7] else ""
        rango = str(row[8]).strip() if row[8] else ""
        data[cr] = {
            "plaza": plaza_corta(str(plaza)),
            "instock": round(instock, 2) if instock is not None else None,
            "clase": clase,
            "rango": rango,
        }
    return data


def leer_visitas(wb):
    """Lee hoja TDAS_SIN_VISITA_DIARIO_DETALLE → dict CR_TIENDA → {tienda, sem_sin_visita, reincidencia}"""
    ws = wb["TDAS_SIN_VISITA_DIARIO_DETALLE"]
    data = {}
    for row in ws.iter_rows(values_only=True):
        plaza = row[5]
        if not plaza or REGION not in str(plaza):
            continue
        cr = str(row[6]).strip()
        tienda = str(row[21]).strip() if row[21] else cr
        sem = int(row[17]) if row[17] is not None else 0
        reinc = str(row[18]).strip() if row[18] else "SIN_REINC"
        data[cr] = {
            "tienda": tienda,
            "sem_sin_visita": sem,
            "reincidencia": reinc,
        }
    return data


def main():
    if len(sys.argv) < 2:
        print("Uso: python convertir.py <archivo.xlsm>")
        sys.exit(1)

    archivo = sys.argv[1]
    print(f"Leyendo: {archivo}")

    wb = load_workbook(archivo, read_only=True, data_only=True)

    fecha = get_fecha(wb)
    fecha_iso = fecha.strftime("%Y-%m-%d")
    fecha_str = fecha.strftime("%-d %b %Y")  # ej: "15 Sep 2026"

    instock = leer_instock(wb)
    visitas = leer_visitas(wb)

    # Combinar ambas fuentes por CR_TIENDA
    todos_cr = set(instock.keys()) | set(visitas.keys())
    tiendas = []
    for cr in todos_cr:
        i = instock.get(cr, {})
        v = visitas.get(cr, {})
        tienda = {
            "cr": cr,
            "tienda": v.get("tienda", cr),
            "plaza": i.get("plaza", v.get("plaza", "")),
            "instock": i.get("instock"),
            "clase": i.get("clase", ""),
            "rango": i.get("rango", ""),
            "sem_sin_visita": v.get("sem_sin_visita", 0),
            "reincidencia": v.get("reincidencia", "SIN_REINC"),
        }
        tiendas.append(tienda)

    # Ordenar: críticas primero, luego por instock ascendente
    tiendas.sort(key=lambda t: (
        REINC_ORDER.get(t["reincidencia"], 99),
        t["instock"] if t["instock"] is not None else 999
    ))

    resultado = {
        "fecha": fecha_iso,
        "corte": fecha_str,
        "tiendas": tiendas
    }

    with open("datos.json", "w", encoding="utf-8") as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2)

    print(f"✅ datos.json generado — {len(tiendas)} tiendas, corte: {fecha_str}")


if __name__ == "__main__":
    main()
